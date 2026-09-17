import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

A = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"           # 本地最新
B = r"cs_pipeline\data\1_Hi-results-12-20260913tclassmate'srevisedversion.xlsx"  # 同学改版

wa = openpyxl.load_workbook(A, read_only=True)
wb = openpyxl.load_workbook(B, read_only=True)
print("本地 sheets:", wa.sheetnames)
print("同学 sheets:", wb.sheetnames)

for sh in wa.sheetnames:
    if sh not in wb.sheetnames:
        print(f"\n### {sh}: 同学版缺此表"); continue
    ra = list(wa[sh].iter_rows(values_only=True))
    rb = list(wb[sh].iter_rows(values_only=True))
    print(f"\n### {sh}: 本地 {len(ra)} 行 vs 同学 {len(rb)} 行")
    if len(ra) != len(rb):
        print("  !! 行数不同")
    diffs = []
    for i in range(max(len(ra), len(rb))):
        rowa = ra[i] if i < len(ra) else None
        rowb = rb[i] if i < len(rb) else None
        if rowa is None or rowb is None:
            diffs.append((i+1, "整行缺失", rowa, rowb)); continue
        for j in range(max(len(rowa), len(rowb))):
            va = rowa[j] if j < len(rowa) else None
            vb = rowb[j] if j < len(rowb) else None
            if va != vb:
                diffs.append((i+1, j+1, va, vb))
    print(f"  差异单元格数: {len(diffs)}")
    for d in diffs[:120]:
        print("   行", d[0], "列", d[1], "| 本地:", repr(d[2]), "| 同学:", repr(d[3]))
    if len(diffs) > 120:
        print(f"   ... 其余 {len(diffs)-120} 处省略")
