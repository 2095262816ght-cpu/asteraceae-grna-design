import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

NOTE = "0913三轮复核：AACATG实际为负链靶点（revcomp=TCTTGAACTCGAACCATGTT+PAM TGG @LG02:147965963），PAM/坐标/链向已更正"

p = r"D:\lifesciencechallenge\cs_pipeline\data\raw\hitom\1_Hi-results-12-20260913_fixed.xlsx"
wb = openpyxl.load_workbook(p)
ws = wb["Sheet1"]
hdr = [c.value for c in ws[1]]
i_seq = hdr.index("seq-20") + 1
i_pam = hdr.index("pam") + 1
i_start = hdr.index("start ") + 1
i_end = hdr.index(" end") + 1
i_strand = hdr.index(" strand") + 1
i_note = hdr.index("备注") + 1
n = 0
for row in ws.iter_rows(min_row=2):
    if row[i_seq - 1].value == "AACATGGTTCGAGTTCAAGA":
        row[i_pam - 1].value = "TGG"
        row[i_start - 1].value = 147965963
        row[i_end - 1].value = 147965983
        row[i_strand - 1].value = "-"
        old = row[i_note - 1].value
        row[i_note - 1].value = (str(old) + "；" if old else "") + NOTE
        n += 1
wb.save(p)
print(f"大表已更新 {n} 行 AACATG -> PAM=TGG, 147965963-147965983, strand=-")

# SY12100 fixed 分表
p2 = r"D:\lifesciencechallenge\cs_pipeline\data\raw\consolidation_sources\all-SY12100-USR-23709-Hi-TOM-sequencing-result_processed_fixed.xlsx"
wb2 = openpyxl.load_workbook(p2)
ws2 = wb2["Sheet1"]
hdr2 = [c.value for c in ws2[1]]
print("SY12100_fixed 列名:", hdr2)
i_seq2 = hdr2.index("seq-20") + 1
i_pam2 = hdr2.index("pam") + 1
i_strand2 = hdr2.index("strand") + 1
i_note2 = hdr2.index("备注") + 1
n2 = 0
for row in ws2.iter_rows(min_row=2):
    if row[i_seq2 - 1].value == "AACATGGTTCGAGTTCAAGA":
        row[i_pam2 - 1].value = "TGG"
        row[i_strand2 - 1].value = "-"
        old = row[i_note2 - 1].value
        row[i_note2 - 1].value = (str(old) + "；" if old else "") + NOTE
        n2 += 1
wb2.save(p2)
print(f"SY12100_fixed 已更新 {n2} 行")
