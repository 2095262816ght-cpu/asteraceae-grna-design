"""D4 打分任务导出缺啥写啥单测：三态标注、百分位口径、缺失原因/所需数据/所用参考包。"""
import io
from pathlib import Path

import pandas as pd
import pytest
from sqlalchemy import delete

from backend.app.db.models import Base, Gene, OfftargetProfileCache, ScoringJob, ScoringJobItem, TargetCandidate
from backend.app.db.session import AsyncSessionLocal, engine
from backend.app.services import export_service
from backend.app.services import offtarget_service
from backend.app.services import reference_bundle_service as rb_service
from backend.app.services.reference_bundle_service import ResolvedBundle

SEQ_PRE = "TTTGTGGATTCCGGTAGAATCTGATGGAAT"   # 将命中预计算候选
SEQ_HIT = "ATATGAAAATGTCAGGATGCCTGAAGGACA"   # 命中基因组（monkeypatch）
SEQ_NONE = "GGGGGGGGGGGGGGGGGGGGGGGGGGGAAA"  # 未锚定
SEQ_BAD = "ACGT"  # 不合法


@pytest.fixture(autouse=True)
async def setup_db():
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    async with AsyncSessionLocal() as session:
        await session.execute(delete(ScoringJobItem))
        await session.execute(delete(ScoringJob))
        await session.execute(delete(TargetCandidate))
        await session.execute(delete(Gene))
        await session.execute(delete(OfftargetProfileCache))
        await session.commit()
    yield


async def _make_job() -> str:
    job = ScoringJob(id="job-export-1", title="导出测试任务", job_type="batch", status="completed",
                     total_sequences=4, processed_sequences=4)
    items = [
        ScoringJobItem(job_id=job.id, custom_id="s1", seq30=SEQ_PRE, pam=SEQ_PRE[24:27],
                       gc_pct=45.0, tm_c=60.0, deepspcas9_score=66.6, crispron_source="wsl_real",
                       efficiency_score=80.0, is_valid=True),
        ScoringJobItem(job_id=job.id, custom_id="s2", seq30=SEQ_HIT, pam=SEQ_HIT[24:27],
                       gc_pct=40.0, tm_c=58.0, deepspcas9_score=55.5, crispron_source="wsl_real",
                       efficiency_score=60.0, is_valid=True),
        ScoringJobItem(job_id=job.id, custom_id="s3", seq30=SEQ_NONE, pam=SEQ_NONE[24:27],
                       gc_pct=90.0, tm_c=70.0, deepspcas9_score=44.4, crispron_source="approximated",
                       efficiency_score=40.0, is_valid=True),
        ScoringJobItem(job_id=job.id, custom_id="s4", seq30=SEQ_BAD, is_valid=False,
                       error_message="序列长度不足 30bp"),
    ]
    async with AsyncSessionLocal() as session:
        session.add(job)
        session.add_all(items)
        session.add(Gene(gene_id="LsT2T1G000010", chrom="Chr1", start=1, end=1000, strand="+",
                         nlr_type="protein-coding", candidate_count=1))
        session.add(TargetCandidate(
            guide_id="T_G_001", gene_id="LsT2T1G000010", chrom="Chr1", start=100, end=130, strand="+",
            seq30=SEQ_PRE, seq20=SEQ_PRE[4:24], pam=SEQ_PRE[24:27], region_type="exon",
            is_centromere=False, is_exon=True, gc_pct=45.0, tm_c=60.0,
            offtarget_count_mis3=3, efficiency_score=82.5, priority_score_leaf=0.61, hazard_leaf=0.12,
            expression_status="mapped", n_intra_panel_sites=1,
        ))
        await session.commit()
    return job.id


@pytest.mark.anyio
async def test_scoring_job_export_honest_columns(monkeypatch):
    job_id = await _make_job()

    bundle = ResolvedBundle(id="mini", display_name="测试参考包", kind="imported", fasta=Path("fake.fa"),
                            gene_source="db", expression_layout="db")

    async def _fake_resolve(db, bundle_id):
        return bundle

    monkeypatch.setattr(rb_service, "resolve_bundle", _fake_resolve)
    # 快导出：不触发现场全基因组扫描；仅复用已按需计算的脱靶缓存（此处预置 s2 的缓存）
    monkeypatch.setattr(
        export_service.offtarget_service, "peek_profiles",
        lambda seq20s, fasta: {
            SEQ_HIT[4:24]: offtarget_service.OfftargetProfile(exact=1, mis1=1, mis2=2, mis3=4)
        },
    )

    async with AsyncSessionLocal() as session:
        content, media_type, filename = await export_service.generate_scoring_job_export_bytes(
            job_id, session, bundle_id="mini", fmt="csv"
        )
    df = pd.read_csv(io.BytesIO(content))

    for col in ("效率百分位", "百分位口径", "多组学证据状态", "缺失原因", "所需数据", "所用参考包"):
        assert col in df.columns

    by_id = {row["样本编号"]: row for _, row in df.iterrows()}

    # 预计算命中：同源复用，无缺失
    assert "命中预计算候选" in by_id["s1"]["多组学证据状态"]
    assert by_id["s1"]["潜在脱靶位点数(<=3)"] in (3, "3", 3.0, "3.0")
    assert float(by_id["s1"]["组织优先指数(叶)"]) == pytest.approx(0.61)
    assert by_id["s1"]["所需数据"] == "—"

    # 已按需算过脱靶（缓存复用）：批内口径百分位 + 真实脱靶计数
    assert "脱靶已按需计算" in by_id["s2"]["多组学证据状态"]
    assert "批内排名口径" in by_id["s2"]["百分位口径"]
    assert by_id["s2"]["潜在脱靶位点数(<=3)"] in (4, "4", 4.0, "4.0")

    # 未命中预计算库且未算过脱靶：快照导出如实标注缺失原因与所需数据
    assert "未命中预计算候选库" in by_id["s3"]["多组学证据状态"]
    assert "未现场执行基因组回贴与全基因组脱靶扫描" in by_id["s3"]["缺失原因"]
    assert "参考基因组回贴" in by_id["s3"]["所需数据"]
    assert "测试参考包" in by_id["s3"]["所用参考包"]

    # 不合法序列：不硬算
    assert "不合法" in by_id["s4"]["多组学证据状态"]
    assert by_id["s4"]["效率百分位"] in ("不适用", "不适用（序列长度不足 30bp）") or "不适用" in str(by_id["s4"]["效率百分位"])


@pytest.mark.anyio
async def test_scoring_job_export_xlsx_has_template_sheet(monkeypatch):
    job_id = await _make_job()

    async def _fake_resolve(db, bundle_id):
        return ResolvedBundle(id="mini", display_name="测试参考包", kind="imported", fasta=None,
                              gene_source="none", expression_layout="none")

    monkeypatch.setattr(rb_service, "resolve_bundle", _fake_resolve)

    async with AsyncSessionLocal() as session:
        content, media_type, filename = await export_service.generate_scoring_job_export_bytes(
            job_id, session, bundle_id="mini", fmt="xlsx"
        )
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "数据表" in wb.sheetnames
    assert "数据需求模板" in wb.sheetnames
    assert "指标方法说明" in wb.sheetnames


@pytest.mark.anyio
async def test_scoring_job_export_reuses_persisted_profile(monkeypatch):
    # 进程内缓存为空（模拟后端重启）：应复用持久化表里的按需计算结果
    job_id = await _make_job()

    bundle = ResolvedBundle(id="mini", display_name="测试参考包", kind="imported", fasta=Path("fake.fa"),
                            gene_source="db", expression_layout="db")

    async def _fake_resolve(db, bundle_id):
        return bundle

    monkeypatch.setattr(rb_service, "resolve_bundle", _fake_resolve)
    monkeypatch.setattr(export_service.offtarget_service, "peek_profiles", lambda seq20s, fasta: {})

    async with AsyncSessionLocal() as session:
        session.add(OfftargetProfileCache(
            genome_key=str(Path("fake.fa")), seq20=SEQ_HIT[4:24],
            exact=1, mis1=1, mis2=2, mis3=7, bundle_id="mini",
        ))
        await session.commit()

    async with AsyncSessionLocal() as session:
        content, media_type, filename = await export_service.generate_scoring_job_export_bytes(
            job_id, session, bundle_id="mini", fmt="csv"
        )
    df = pd.read_csv(io.BytesIO(content))
    by_id = {row["样本编号"]: row for _, row in df.iterrows()}

    assert "脱靶已按需计算" in by_id["s2"]["多组学证据状态"]
    assert by_id["s2"]["潜在脱靶位点数(<=3)"] in (7, "7", 7.0, "7.0")
    # 未算过的行照旧如实标注，不以 0 冒充
    assert "未命中预计算候选库" in by_id["s3"]["多组学证据状态"]
    assert by_id["s3"]["潜在脱靶位点数(<=3)"] in ("未计算",)


@pytest.mark.anyio
async def test_scoring_job_export_falls_back_to_job_bundle(monkeypatch):
    # 历史导出不带 bundle_id 时沿用任务记录的参考包，不再固定 lettuce
    job_id = await _make_job()
    seen = {}

    async def _fake_resolve(db, bundle_id):
        seen["bundle_id"] = bundle_id
        return ResolvedBundle(id=bundle_id, display_name="测试参考包", kind="imported", fasta=None,
                              gene_source="none", expression_layout="none")

    monkeypatch.setattr(rb_service, "resolve_bundle", _fake_resolve)

    async with AsyncSessionLocal() as session:
        job = await session.get(ScoringJob, job_id)
        job.bundle_id = "tks"
        await session.commit()

    async with AsyncSessionLocal() as session:
        await export_service.generate_scoring_job_export_bytes(job_id, session, bundle_id=None, fmt="csv")
    assert seen["bundle_id"] == "tks"
