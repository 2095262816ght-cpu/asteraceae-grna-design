import sys, io, shutil, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

XLSX = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"

ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = XLSX.replace(".xlsx", f"_backup_{ts}_precattcoordfix.xlsx")
shutil.copy2(XLSX, bak)
print("备份:", bak)

wb = openpyxl.load_workbook(XLSX)
ws = wb["Sheet1"]
hdr = [str(c.value).strip() for c in ws[1]]
c = {h: i+1 for i, h in enumerate(hdr)}

n = 0
for row in range(2, ws.max_row+1):
    if str(ws.cell(row, c["seq-20"]).value).strip().upper() == "CATTACCAAAATAAAACAGG" \
       and str(ws.cell(row, c["chrom"]).value).strip() == "LG07" \
       and ws.cell(row, c["start"]).value == 109029137:
        assert ws.cell(row, c["end"]).value == 109029157
        ws.cell(row, c["start"]).value = 109029136
        ws.cell(row, c["end"]).value = 109029156
        n += 1
wb.save(XLSX)
print(f"改动 {n} 行（应=13）：start 109029137→109029136, end 109029157→109029156")
assert n == 13
print("断言通过")
