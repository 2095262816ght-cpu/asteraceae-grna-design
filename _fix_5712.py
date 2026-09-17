import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

PATH = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Sheet1"]

hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
col = {name: i+1 for i, name in enumerate(hdr)}
print("表头:", hdr)

NOTE = ("原记录seq-20/pam/坐标有误(GAAAATATTGAAAACATATA+TGG@LG07:109029054-074),"
        "经原始reads核实实际编辑位点更正为GAAAATGTCAGGATGCCTGA+AGG@LG07:109028897-917(负链),"
        "详见docs/problem.md第4项(2026-09-15)")

n = 0
for row in range(2, ws.max_row+1):
    if str(ws.cell(row, col["batch"]).value).strip() == "USR-5712":
        ws.cell(row, col["seq-20"]).value = "GAAAATGTCAGGATGCCTGA"
        ws.cell(row, col["pam"]).value = "AGG"
        ws.cell(row, col["start"]).value = 109028897
        ws.cell(row, col["end"]).value = 109028917
        ws.cell(row, col["备注"]).value = NOTE
        n += 1
wb.save(PATH)
print(f"已改正 {n} 行")
