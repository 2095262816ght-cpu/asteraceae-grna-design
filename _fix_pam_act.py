import sys, io, shutil, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

XLSX = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"
GUIDE = "AACATGGTTCGAGTTCAAGA"
TARGETS = {"TKS_0260", "TKS_0261", "TKS_0262", "TKS_0264"}
NOTE = ("PAM原记CCA，基因组实测为ACT（2026-09-16改正）；该位点无NGG PAM，"
        "实测基本无编辑活性（0.0473为已知天然SNP背景），属无效设计；"
        "00脚本已按非NGG过滤剔除，不影响训练")

genome = {}
cur = None
for line in open(r"cs_pipeline\data\reference_genomes\2-1_TKS.genome.fasta"):
    if line.startswith(">"):
        cur = line[1:].split()[0]; genome[cur] = []
    elif cur:
        genome[cur].append(line.strip().upper())
genome = {k: "".join(v) for k, v in genome.items()}

# 1) 备份
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = XLSX.replace(".xlsx", f"_backup_{ts}_prepamactfix.xlsx")
shutil.copy2(XLSX, bak)
print("备份:", bak)

# 2) 基因组侧验证：guide 下游 3bp 必须确为 ACT
gseq = genome["LG02"]
pos = gseq.find(GUIDE)
assert pos != -1 and gseq.find(GUIDE, pos + 1) == -1, "guide 在 LG02 应唯一命中"
downstream = gseq[pos + 20:pos + 23]
assert downstream == "ACT", f"基因组下游 3bp 应为 ACT，实测 {downstream}"
print(f"基因组验证通过：LG02:{pos+1}-{pos+20} 下游 3bp = {downstream}")

# 3) 改正 4 行
wb = openpyxl.load_workbook(XLSX)
ws = wb["Sheet1"]
hdr = [str(c.value).strip() for c in ws[1]]
c = {h: i + 1 for i, h in enumerate(hdr)}

changed = []
for row in range(2, ws.max_row + 1):
    gid = ws.cell(row, c["guide-id"]).value
    if not gid or str(gid).strip() not in TARGETS:
        continue
    s = str(ws.cell(row, c["seq-20"]).value).strip().upper()
    assert s == GUIDE, f"{gid} seq-20 不符预期: {s}"
    pam_cell = ws.cell(row, c["pam"])
    old_pam = str(pam_cell.value).strip()
    assert old_pam == "CCA", f"{gid} pam 应为 CCA，实际 {old_pam}"
    pam_cell.value = "ACT"
    note_cell = ws.cell(row, c["备注"])
    old_note = note_cell.value
    note_cell.value = (str(old_note).strip() + "；" + NOTE) if old_note else NOTE
    changed.append((str(gid).strip(), old_pam, "ACT"))

wb.save(XLSX)
print("改动行:", changed)
assert len(changed) == 4 and {g for g, _, _ in changed} == TARGETS, "改动行数/对象不符预期!"
print("断言通过：4 行 pam CCA -> ACT，备注已追加")

# 4) 读回验证
wb2 = openpyxl.load_workbook(XLSX)
ws2 = wb2["Sheet1"]
n = 0
for row in range(2, ws2.max_row + 1):
    gid = ws2.cell(row, c["guide-id"]).value
    if gid and str(gid).strip() in TARGETS:
        assert str(ws2.cell(row, c["pam"]).value).strip() == "ACT"
        n += 1
assert n == 4
print("读回验证通过：4 行 pam 均为 ACT")
