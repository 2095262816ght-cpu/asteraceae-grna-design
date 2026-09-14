"""逐项核对 problem.md 中的问题在 data\\1_Hi-results-12-20260913.xlsx 当前状态"""
import openpyxl
from collections import Counter

PATH = r"d:\lifesciencechallenge\cs_pipeline\data\1_Hi-results-12-20260913.xlsx"
wb = openpyxl.load_workbook(PATH)
print("Sheets:", wb.sheetnames)

ws = wb["Sheet1"]
rows = list(ws.iter_rows(values_only=True))
header = rows[0]
print("Header:", [repr(h) for h in header])
data = rows[1:]
print("Data rows:", len(data))

# 找列索引（容错表头空格）
def col(name):
    for i, h in enumerate(header):
        if h and str(h).strip().lower() == name.lower():
            return i
    return None

c_uid = col("guide-id")
c_gene = col("gene-id")
c_seq = col("seq-20")
c_pam = col("PAM")
c_chrom = col("chrom")
c_start = col("start")
c_end = col("end")
c_strand = col("strand")
c_batch = col("batch")
c_rep = col("rep")
c_indel = col("indel_pct")
c_note = col("备注") if col("备注") is not None else col("note")
print("cols:", dict(uid=c_uid, gene=c_gene, seq=c_seq, pam=c_pam, chrom=c_chrom,
                    start=c_start, end=c_end, strand=c_strand, batch=c_batch,
                    rep=c_rep, indel=c_indel, note=c_note))

# --- 问题1: TKS_0404 前导空格 & indel=1 & USR-8146 ---
print("\n=== 问题1: TKS_0404 ===")
for r in data:
    if r[c_uid] == "TKS_0404":
        print("  seq:", repr(r[c_seq]), "| pam:", r[c_pam], "| indel:", r[c_indel],
              "| batch:", repr(r[c_batch]), "| gene:", repr(r[c_gene]),
              "| start:", r[c_start], "| end:", r[c_end], "| note:", repr(r[c_note]))
# 其他前导/尾随空格检查
bad_seq = [r[c_uid] for r in data if r[c_seq] and str(r[c_seq]) != str(r[c_seq]).strip()]
print("  其他含首尾空格的 seq:", bad_seq)
# USR-8146
b8146 = [r[c_uid] for r in data if str(r[c_batch]).strip() == "USR-8146"]
print("  USR-8146 行:", b8146)

# --- 问题2: TKS_0263 编号错位 ---
print("\n=== 问题2: TKS_0260~0264 ===")
for r in data:
    if r[c_uid] in ("TKS_0260","TKS_0261","TKS_0262","TKS_0263","TKS_0264"):
        print(" ", r[c_uid], "| gene:", repr(r[c_gene]), "| seq:", r[c_seq],
              "| pam:", r[c_pam], "| pos:", r[c_chrom], r[c_start], r[c_end],
              "| batch:", r[c_batch], "| indel:", r[c_indel], "| note:", repr(r[c_note]))

# --- 问题3: 4 行 CCA PAM ---
print("\n=== 问题3: CCA PAM 行 ===")
pam_counter = Counter(str(r[c_pam]).strip() for r in data)
print("  PAM 分布:", dict(pam_counter))
for r in data:
    if str(r[c_pam]).strip() == "CCA":
        print(" ", r[c_uid], "| gene:", repr(r[c_gene]), "| seq:", r[c_seq],
              "| pos:", r[c_chrom], r[c_start], r[c_end], "| batch:", r[c_batch],
              "| indel:", r[c_indel], "| note:", repr(r[c_note]))

# --- 问题4: strand 列 ---
print("\n=== 问题4: strand 分布 ===")
print("  ", dict(Counter(str(r[c_strand]).strip() for r in data)))

# --- 问题5: efficiency-norm ---
c_eff = col("efficiency-norm")
print("\n=== 问题5: efficiency-norm ===", "列索引:", c_eff)
if c_eff is not None:
    vals = [r[c_eff] for r in data]
    nonempty = [v for v in vals if v not in (None, "")]
    print("  非空行数:", len(nonempty), "/", len(vals))

# --- 问题6: 表头空格 / 第一列 / rep 类型 ---
print("\n=== 问题6: 表头 ===")
print("  带空格表头:", [repr(h) for h in header if h and str(h) != str(h).strip()])
print("  第一列表头:", repr(header[0]), "非空值数:", sum(1 for r in data if r[0] not in (None,"")))
rep_types = Counter(type(r[c_rep]).__name__ for r in data)
rep_vals = Counter(str(r[c_rep]) for r in data)
print("  rep 类型:", dict(rep_types))
print("  rep 值分布:", dict(rep_vals.most_common(10)))

# --- 问题7: 辅助 sheet ---
print("\n=== 问题7: 辅助 sheet ===")
for name in wb.sheetnames:
    s = wb[name]
    rs = list(s.iter_rows(values_only=True))
    print(f"  {name}: {len(rs)} 行 x {s.max_column} 列; 首行: {[repr(x) for x in rs[0]][:12]}")
    if name in ("Sheet2","Sheet3","Sheet4","Sheet5"):
        for r in rs[1:4]:
            print("    ", [repr(x) for x in r][:12])

# --- 补充问题: 60 行 100% ---
print("\n=== 补充3: indel=1 的行 ===")
ones = [(r[c_uid], r[c_gene], r[c_batch]) for r in data if r[c_indel] == 1]
print("  数量:", len(ones))
print("  分布:", dict(Counter((g, b) for _, g, b in ones)))

# --- 补充5: USR-5712 编辑率 ---
print("\n=== 补充5: USR-5712 ===")
b5712 = [r[c_indel] for r in data if str(r[c_batch]).strip() == "USR-5712"]
print("  行数:", len(b5712), "indel 值集合:", sorted(set(b5712)))

# --- 补充6: TKS_0160 ---
print("\n=== 补充6: TKS_0160 ===")
for r in data:
    if r[c_uid] == "TKS_0160":
        print("  indel:", r[c_indel], "| note:", repr(r[c_note]))

# --- TKS_0184 ---
print("\n=== TKS_0184 (离群 0.5151) ===")
for r in data:
    if r[c_uid] == "TKS_0184":
        print("  indel:", r[c_indel], "| note:", repr(r[c_note]))

# --- TKS_0151/0185/0213 口径修正值 ---
print("\n=== 口径修正三行 ===")
for r in data:
    if r[c_uid] in ("TKS_0151","TKS_0185","TKS_0213"):
        print(" ", r[c_uid], "indel:", r[c_indel], "| note:", repr(r[c_note]))

# 空 indel 行
empty_indel = [r[c_uid] for r in data if r[c_indel] in (None, "")]
print("\nindel 为空的行:", empty_indel)

# gene-id 尾随空格
bad_gene = [r[c_uid] for r in data if r[c_gene] and str(r[c_gene]) != str(r[c_gene]).strip()]
print("gene-id 含首尾空格的行数:", len(bad_gene), bad_gene[:10])
