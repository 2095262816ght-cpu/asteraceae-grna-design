import openpyxl, shutil, os, datetime
from pathlib import Path

SRC = str(Path(__file__).resolve().parents[1] / "data" / "1_Hi-results-12-20260913.xlsx")

# 备份
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = SRC.replace(".xlsx", f"_backup_{ts}.xlsx")
shutil.copy2(SRC, bak)
print(f"备份 -> {bak}")

wb = openpyxl.load_workbook(SRC)
ws = wb["Sheet1"]

hdr = [c.value for c in ws[1]]
def col_idx(name):
    for i, h in enumerate(hdr):
        if h is not None and str(h).strip().lower() == name.lower():
            return i
    return None

ci_guide = col_idx("guide-id")
ci_seq   = col_idx("seq-20")
ci_pam   = col_idx("pam")
ci_batch = col_idx("batch")
ci_note  = col_idx("备注")
print("列索引:", ci_guide, ci_seq, ci_pam, ci_batch, ci_note)

fixed = {"seq_space": [], "batch": [], "header": []}

# 1) 表头空格 strip
for i, h in enumerate(hdr):
    if h is not None and isinstance(h, str) and h != h.strip():
        ws.cell(row=1, column=i+1, value=h.strip())
        fixed["header"].append(f"{h!r} -> {h.strip()!r}")

# 2) 数据行修复
for r in range(2, ws.max_row + 1):
    gid = ws.cell(row=r, column=ci_guide+1).value
    seq = ws.cell(row=r, column=ci_seq+1).value
    bat = ws.cell(row=r, column=ci_batch+1).value

    # seq 前导/尾随空格
    if isinstance(seq, str) and seq != seq.strip():
        ws.cell(row=r, column=ci_seq+1, value=seq.strip())
        fixed["seq_space"].append((r, gid, seq.strip()))

    # USR-8146 -> USR-8145（原始文件 USR-8145 块号 097-192 无 096，全表仅此一行）
    if isinstance(bat, str) and bat.strip() == "USR-8146":
        ws.cell(row=r, column=ci_batch+1, value="USR-8145")
        note = ws.cell(row=r, column=ci_note+1).value
        add = "批次号由USR-8146更正为USR-8145(原文件块号097-192无096块)"
        ws.cell(row=r, column=ci_note+1,
                value=(str(note).strip() + ";" + add) if note else add)
        fixed["batch"].append((r, gid))

print("\n=== 表头修复 ===")
for x in fixed["header"]: print(" ", x)
print(f"\n=== seq 空格修复 {len(fixed['seq_space'])} 行 ===")
for r, g, s in fixed["seq_space"]:
    print(f"  行{r} {g}: {s}")
print(f"\n=== 批次修复 {len(fixed['batch'])} 行 ===")
for r, g in fixed["batch"]: print(f"  行{r} {g}: USR-8146 -> USR-8145")

wb.save(SRC)
print("\n已保存:", SRC)
