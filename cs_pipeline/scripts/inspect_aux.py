import openpyxl
wb = openpyxl.load_workbook(r"d:\lifesciencechallenge\cs_pipeline\data\1_Hi-results-12-20260913.xlsx", data_only=False)

for sn in ["Sheet2", "Sheet3", "Sheet4", "Sheet5"]:
    ws = wb[sn]
    print(f"\n=== {sn} ({ws.max_row}x{ws.max_column}) ===")
    for r in range(1, min(ws.max_row, 6) + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"  行{r}: {row}")
