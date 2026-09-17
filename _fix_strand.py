import sys, io, shutil, datetime
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

XLSX = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"

def revcomp(s): return s.translate(str.maketrans("ACGT","TGCA"))[::-1]

genome={}; cur=None
for line in open(r"cs_pipeline\data\reference_genomes\2-1_TKS.genome.fasta"):
    if line.startswith(">"): cur=line[1:].split()[0]; genome[cur]=[]
    elif cur: genome[cur].append(line.strip().upper())
genome={k:"".join(v) for k,v in genome.items()}

# 1) 备份
ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
bak = XLSX.replace(".xlsx", f"_backup_{ts}_prestrandfix.xlsx")
shutil.copy2(XLSX, bak)
print("备份:", bak)

# 2) 打开并逐行判定
wb = openpyxl.load_workbook(XLSX)
ws = wb["Sheet1"]
hdr = [str(c.value).strip() for c in ws[1]]
c = {h: i+1 for i, h in enumerate(hdr)}  # 1-based 列号

n_plus = n_minus = n_skip = 0
changed = []
for row in range(2, ws.max_row+1):
    gid = ws.cell(row, c["guide-id"]).value
    if not gid: continue
    s = str(ws.cell(row, c["seq-20"]).value).strip().upper()
    if len(s)!=20 or set(s)-set("ACGT"):
        n_skip += 1; continue
    chrom = str(ws.cell(row, c["chrom"]).value).strip()
    gseq = genome.get(chrom, "")
    if s in gseq:
        new = "+"
    elif revcomp(s) in gseq:
        new = "-"
    else:
        print("!! 基因组未命中，跳过:", gid, s); n_skip += 1; continue
    old = ws.cell(row, c["strand"]).value
    if str(old).strip() != new:
        ws.cell(row, c["strand"]).value = new
        changed.append((gid, old, new))
    if new == "+": n_plus += 1
    else: n_minus += 1

wb.save(XLSX)
print(f"正链(+) {n_plus} 行 / 负链(-) {n_minus} 行 / 跳过 {n_skip} 行")
print(f"实际改动 {len(changed)} 行（应=221，全部 - → +）")
assert len(changed)==221 and all(o=="-" and n=="+" for _,o,n in changed), "改动不符合预期!"
print("断言通过：221 行 - → +，无其他改动")
