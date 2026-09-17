import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import openpyxl

PATH = r"cs_pipeline\data\1_Hi-results-12-20260913.xlsx"
wb = openpyxl.load_workbook(PATH)
ws = wb["Sheet1"]
hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
col = {name: i+1 for i, name in enumerate(hdr)}

NOTE = ("原记录seq-20/pam/坐标为移位5bp的错误窗口(GGATTGCTGCTGAGAAGGCC+TGG,"
        "基因组实测PAM位为ATG不合法),经原始reads切割点定位及Sheet2记录核实,"
        "更正为CGTAAGGATTGCTGCTGAGA+AGG(负链),详见docs/problem.md第2项(2026-09-15)")

CO = {"evm.model.LG01.6995": (165462914, 165462934),
      "evm.model.LG01.6999": (165502832, 165502852)}
n = 0
for row in range(2, ws.max_row+1):
    if str(ws.cell(row, col["seq-20"]).value).strip() == "GGATTGCTGCTGAGAAGGCC":
        gene = str(ws.cell(row, col["gene-id"]).value).strip()
        st, en = CO[gene]
        ws.cell(row, col["seq-20"]).value = "CGTAAGGATTGCTGCTGAGA"
        ws.cell(row, col["pam"]).value = "AGG"
        ws.cell(row, col["start"]).value = st
        ws.cell(row, col["end"]).value = en
        ws.cell(row, col["备注"]).value = NOTE
        n += 1
wb.save(PATH)
print(f"已改正 {n} 行")

# 验证
import pandas as pd
df = pd.ExcelFile(PATH).parse("Sheet1")
df.columns = [str(c).strip() for c in df.columns]
print("旧序列残留:", (df["seq-20"].astype(str).str.strip()=="GGATTGCTGCTGAGAAGGCC").sum())
sub = df[df["seq-20"].astype(str).str.strip()=="CGTAAGGATTGCTGCTGAGA"]
for k, g in sub.groupby(["gene-id","pam","chrom","start","end","strand"]):
    print(" ", k, f"x{len(g)}")
